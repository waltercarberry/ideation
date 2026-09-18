"""Prompt layer per artifact -> one Doc model -> Word (.docx) and standalone HTML.

The LLM only ever sees exact, pre-computed FACTS and is told not to calculate. If it is
unavailable, each artifact still builds from deterministic sentences, so exports never fail.
"""
import base64
import html
import io
import json
from dataclasses import dataclass, field
from datetime import datetime
from typing import Literal

import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from pydantic import BaseModel

import llm
from charts import BLUE, INK, PINK, YELLOW, build_frame, digest, mpl_png, plan_charts
from data_processor import DatasetProfile, compute_facts, fmt, label

KINDS = {"monthly": "Monthly report", "weekly": "Weekly insights", "visuals": "Data visualisations"}
SYSTEM = ("You are a senior data analyst writing for busy executives. Use ONLY the numbers in FACTS: never calculate, "
          "estimate or invent a figure. Plain, direct sentences, no hype. Use the human-readable 'label' of each metric.")


# ------------------------------------------------------------------ document model
@dataclass
class Doc:
    kind: str
    title: str
    subtitle: str = ""
    headline: str = ""
    kpis: list = field(default_factory=list)    # (label, value, delta)
    blocks: list = field(default_factory=list)  # ("h2", t) ("p", t) ("bullets", [t]) ("insight", title, detail, dir) ("chart", spec, frame, caption)
    notes: list = field(default_factory=list)
    ai: bool = True


class Section(BaseModel):
    heading: str
    body: str


class MonthlyNarrative(BaseModel):
    headline: str
    sections: list[Section]
    recommendations: list[str]


class Insight(BaseModel):
    title: str
    detail: str
    direction: Literal["up", "down", "flat", "watch"]


class WeeklyNarrative(BaseModel):
    headline: str
    insights: list[Insight]
    watchlist: list[str]


class Captions(BaseModel):
    captions: list[str]


# ------------------------------------------------------------------ deterministic fallbacks
def _metric_lines(facts: dict) -> list[str]:
    out = []
    for m in facts["metrics"]:
        u, name = m["unit"], m["label"]
        if m.get("change_pct") is not None:
            way = "up" if m["change_pct"] > 0 else "down" if m["change_pct"] < 0 else "flat"
            word = "total" if m["aggregation"] == "sum" else "average"
            out.append(f"{name} {word} was {fmt(m['current_period_value'], u)}, {way} {abs(m['change_pct'])}% on the previous period ({fmt(m['previous_period_value'], u)}).")
        else:
            out.append(f"{name} averaged {fmt(m['overall']['mean'], u)} (range {fmt(m['overall']['min'], u)} to {fmt(m['overall']['max'], u)}).")
    return out


def _category_lines(facts: dict) -> list[str]:
    out = []
    for m in facts["metrics"]:
        c = m.get("by_category")
        if c and c["highest"]:
            hi, lo = c["highest"][0], c["lowest"][0]
            out.append(f"By {label(c['column']).lower()}, {hi[0]} leads on {m['label'].lower()} ({fmt(hi[1], m['unit'])}) and {lo[0]} trails ({fmt(lo[1], m['unit'])}).")
    return out


def _fallback_monthly(facts):
    lines = _metric_lines(facts)
    return MonthlyNarrative(
        headline=lines[0] if lines else facts["dataset"],
        sections=[Section(heading="What moved", body=" ".join(lines) or "No numeric measures were found."),
                  Section(heading="What's driving it", body=" ".join(_category_lines(facts)) or "No categories were available to break the numbers down."),
                  Section(heading="What to watch", body=" ".join(facts["data_quality_notes"]) or "No data quality issues were found.")],
        recommendations=[])


def _fallback_weekly(facts):
    ins = [Insight(title=m["label"], detail=l, direction="up" if (m.get("change_pct") or 0) > 0 else "down" if (m.get("change_pct") or 0) < 0 else "flat")
           for m, l in zip(facts["metrics"], _metric_lines(facts))]
    return WeeklyNarrative(headline=ins[0].detail if ins else facts["dataset"], insights=ins, watchlist=facts["data_quality_notes"][:3])


# ------------------------------------------------------------------ prompt layers
def _monthly(facts):
    w = facts.get("window", {})
    prompt = f"""FACTS:
{json.dumps(facts, indent=1)}

Write a monthly performance report for the period {w.get('current_period', facts.get('date_range', ''))}, compared with the previous period where FACTS allows.
- headline: one sentence with the single most important thing that happened.
- sections: exactly three, with headings "What moved", "What's driving it" and "What to watch". One or two short paragraphs each.
- recommendations: three concrete actions, one sentence each, each tied to a number in FACTS."""
    try:
        return llm.ask_json(prompt, MonthlyNarrative, SYSTEM), True
    except llm.LLMUnavailable:
        return _fallback_monthly(facts), False


def _weekly(facts):
    w = facts.get("window", {})
    prompt = f"""FACTS:
{json.dumps(facts, indent=1)}

Write short weekly insights for {w.get('current_period', facts.get('date_range', ''))}.
- headline: one sentence.
- insights: 3 to 5 items. title (max 8 words), detail (one or two sentences quoting numbers from FACTS), direction (up, down, flat or watch).
- watchlist: up to three specific things to check next week."""
    try:
        return llm.ask_json(prompt, WeeklyNarrative, SYSTEM), True
    except llm.LLMUnavailable:
        return _fallback_weekly(facts), False


def _captions(digests: list[str]) -> tuple[list[str], bool]:
    if not digests:
        return [], True
    prompt = f"""Each line describes one chart with exact figures:
{json.dumps(digests, indent=1)}

Return one caption per chart, in the same order. Each caption is one sentence saying what the reader should take away,
keeping the figures exactly as given."""
    try:
        caps = [c.strip() for c in llm.ask_json(prompt, Captions, SYSTEM).captions]
        return (caps + digests[len(caps):])[:len(digests)], True
    except llm.LLMUnavailable:
        return digests, False


# ------------------------------------------------------------------ build
def _kpis(facts: dict) -> list:
    out = []
    for m in facts["metrics"]:
        val = m.get("current_period_value")
        val = m["overall"]["total" if m["aggregation"] == "sum" else "mean"] if val is None else val
        chg = m.get("change_pct")
        out.append((m["label"], fmt(val, m["unit"]), f"{chg:+.1f}% vs previous period" if chg is not None else ""))
    return out


def build_document(kind: str, df: pd.DataFrame, profile: DatasetProfile, audit_log: list[str]) -> Doc:
    facts = compute_facts(df, profile, kind, audit_log)
    doc = Doc(kind=kind, title=KINDS[kind], notes=audit_log, subtitle=profile.summary)
    if kind != "visuals":
        doc.subtitle = f"{profile.summary} Period: {facts.get('window', {}).get('current_period') or facts.get('date_range', 'all data')}."
        doc.kpis = _kpis(facts)

    n_charts = {"monthly": 2, "weekly": 1, "visuals": 5}[kind]
    charts = [(s, build_frame(df, s)) for s in plan_charts(df, profile, n_charts)]
    units = {c.name: c.unit for c in profile.columns}
    caps, ai_caps = _captions([digest(f, s, units) for s, f in charts])
    chart_blocks = [("chart", s, f, c) for (s, f), c in zip(charts, caps)]

    if kind == "monthly":
        n, ai = _monthly(facts)
        doc.headline = n.headline
        for s in n.sections:
            doc.blocks += [("h2", s.heading), ("p", s.body)]
        if chart_blocks:
            doc.blocks += [("h2", "The numbers")] + chart_blocks
        if n.recommendations:
            doc.blocks += [("h2", "Recommended actions"), ("bullets", n.recommendations)]
    elif kind == "weekly":
        n, ai = _weekly(facts)
        doc.headline = n.headline
        doc.blocks += [("insight", i.title, i.detail, i.direction) for i in n.insights] + chart_blocks
        if n.watchlist:
            doc.blocks += [("h2", "Watch next week"), ("bullets", n.watchlist)]
    else:
        ai = ai_caps
        doc.headline = profile.summary
        doc.subtitle = f"{len(charts)} charts chosen for this data."
        doc.blocks += chart_blocks
    doc.ai = ai and ai_caps
    return doc


# ------------------------------------------------------------------ Word
def _hex(h):
    return RGBColor.from_string(h.lstrip("#"))


def _shade(cell, hex_):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_.lstrip("#"))
    cell._tc.get_or_add_tcPr().append(shd)


def _style_font(style, size, bold=False, color=INK):
    style.font.name = "Helvetica"
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = _hex(color)
    rf = style.element.get_or_add_rPr().get_or_add_rFonts()
    for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        rf.attrib.pop(qn(a), None)
    for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rf.set(qn(a), "Helvetica")


def _run(p, text, size=None, bold=False, color=INK, italic=False):
    r = p.add_run(text)
    r.bold, r.italic = bold, italic
    r.font.color.rgb = _hex(color)
    if size:
        r.font.size = Pt(size)
    return r


def to_docx(doc: Doc) -> bytes:
    d = Document()
    sec = d.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)  # A4
    sec.left_margin = sec.right_margin = Cm(2)
    sec.top_margin = sec.bottom_margin = Cm(2)
    _style_font(d.styles["Normal"], 10.5)
    d.styles["Normal"].paragraph_format.space_after = Pt(6)
    _style_font(d.styles["Heading 1"], 26, True)
    _style_font(d.styles["Heading 2"], 15, True, BLUE)
    _style_font(d.styles["List Bullet"], 10.5)

    # title block
    t = d.add_table(rows=1, cols=1)
    c = t.rows[0].cells[0]
    _shade(c, YELLOW)
    p = c.paragraphs[0]
    p.paragraph_format.space_before = Pt(14)
    _run(p, doc.title, 30, True)
    p2 = c.add_paragraph()
    p2.paragraph_format.space_after = Pt(14)
    _run(p2, doc.subtitle, 10)
    if doc.headline:
        p = d.add_paragraph()
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(10)
        _run(p, doc.headline, 16, True)

    if doc.kpis:
        t = d.add_table(rows=1, cols=len(doc.kpis))
        for cell, (lab, val, delta), fill in zip(t.rows[0].cells, doc.kpis, [YELLOW, PINK, BLUE, YELLOW]):
            _shade(cell, fill)
            ink = "#FFFFFF" if fill == BLUE else INK
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(8)
            _run(p, lab, 9, color=ink)
            p = cell.add_paragraph()
            p.paragraph_format.space_after = Pt(0)
            _run(p, val, 20, True, ink)
            p = cell.add_paragraph()
            p.paragraph_format.space_after = Pt(8)
            _run(p, delta or " ", 8.5, color=ink)
        d.add_paragraph()

    dir_word = {"up": ("Up", BLUE), "down": ("Down", PINK), "flat": ("Steady", INK), "watch": ("Watch", INK)}
    for b in doc.blocks:
        if b[0] == "h2":
            d.add_heading(b[1], level=2)
        elif b[0] == "p":
            d.add_paragraph(b[1])
        elif b[0] == "bullets":
            for item in b[1]:
                d.add_paragraph(item, style="List Bullet")
        elif b[0] == "insight":
            word, col = dir_word[b[3]]
            p = d.add_paragraph()
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.keep_with_next = True
            _run(p, f"{word}  ", 10.5, True, col)
            _run(p, b[1], 12, True)
            d.add_paragraph(b[2])
        elif b[0] == "chart":
            d.add_heading(b[1].title, level=2).paragraph_format.keep_with_next = True
            d.add_picture(io.BytesIO(mpl_png(b[2], b[1])), width=Cm(16.5))
            d.paragraphs[-1].paragraph_format.keep_with_next = True
            p = d.add_paragraph()
            _run(p, b[3], 9.5, italic=True)

    if doc.notes:
        d.add_heading("Data quality", level=2)
        for n in doc.notes:
            d.add_paragraph(n, style="List Bullet")
    p = d.add_paragraph()
    p.paragraph_format.space_before = Pt(18)
    how = f"written by {llm.MODEL}" if doc.ai else "written with built-in rules (the AI model was unavailable)"
    _run(p, f"Generated by AutoReport AI on {datetime.now():%d %b %Y}. All figures are calculated from the cleaned data; the narrative is {how}.", 8, color=INK)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


# ------------------------------------------------------------------ HTML
_CSS = f"""
:root {{ --ink:{INK}; --blue:{BLUE}; --pink:{PINK}; --yellow:{YELLOW}; --paper:#F3F3EE; }}
* {{ box-sizing:border-box; }}
body {{ margin:0; background:var(--paper); color:var(--ink); font:18px/1.5 "Helvetica Neue",Helvetica,Arial,sans-serif; }}
main {{ max-width:860px; margin:0 auto; padding:3rem 1.5rem 5rem; }}
header {{ background:var(--yellow); border:3px solid var(--ink); box-shadow:10px 10px 0 var(--pink); padding:2rem 2.2rem; }}
h1 {{ font-size:clamp(2.4rem,7vw,4.2rem); line-height:.95; letter-spacing:-.05em; margin:0 0 .6rem; }}
h2 {{ font-size:1.9rem; letter-spacing:-.03em; margin:3rem 0 .6rem; color:var(--blue); }}
.sub {{ margin:0; font-size:1rem; }}
.headline {{ font-size:1.7rem; font-weight:700; letter-spacing:-.03em; line-height:1.15; margin:2.5rem 0 0; }}
.kpis {{ display:grid; grid-template-columns:repeat(auto-fit,minmax(180px,1fr)); gap:1rem; margin-top:2rem; }}
.kpi {{ border:3px solid var(--ink); border-radius:18px; padding:1rem 1.2rem; background:var(--yellow); }}
.kpi:nth-child(3n+2) {{ background:var(--pink); }} .kpi:nth-child(3n) {{ background:var(--blue); color:#fff; }}
.kpi b {{ display:block; font-size:2.2rem; letter-spacing:-.04em; line-height:1.1; }} .kpi small {{ font-size:.85rem; }}
.insight {{ margin:1.6rem 0 0; }} .insight strong {{ font-size:1.2rem; }}
.tag {{ display:inline-block; border:2.5px solid var(--ink); border-radius:999px; padding:0 .8rem; font-weight:700; font-size:.9rem; margin-right:.6rem; background:var(--paper); }}
.tag.up {{ background:var(--blue); color:#fff; }} .tag.down {{ background:var(--pink); }}
figure {{ margin:1.5rem 0; }} figure img {{ width:100%; border:3px solid var(--ink); box-shadow:8px 8px 0 var(--blue); background:#fff; }}
figcaption {{ margin-top:1rem; font-size:.95rem; font-style:italic; }}
footer {{ margin-top:4rem; font-size:.8rem; }}
"""


def to_html(doc: Doc) -> str:
    e = html.escape
    parts = [f"<header><h1>{e(doc.title)}</h1><p class='sub'>{e(doc.subtitle)}</p></header>"]
    if doc.headline:
        parts.append(f"<p class='headline'>{e(doc.headline)}</p>")
    if doc.kpis:
        parts.append("<div class='kpis'>" + "".join(
            f"<div class='kpi'>{e(l)}<b>{e(v)}</b><small>{e(dl) or '&nbsp;'}</small></div>" for l, v, dl in doc.kpis) + "</div>")
    words = {"up": "Up", "down": "Down", "flat": "Steady", "watch": "Watch"}
    for b in doc.blocks:
        if b[0] == "h2":
            parts.append(f"<h2>{e(b[1])}</h2>")
        elif b[0] == "p":
            parts.append(f"<p>{e(b[1])}</p>")
        elif b[0] == "bullets":
            parts.append("<ul>" + "".join(f"<li>{e(i)}</li>" for i in b[1]) + "</ul>")
        elif b[0] == "insight":
            parts.append(f"<div class='insight'><span class='tag {b[3]}'>{words[b[3]]}</span><strong>{e(b[1])}</strong><p>{e(b[2])}</p></div>")
        elif b[0] == "chart":
            img = base64.b64encode(mpl_png(b[2], b[1])).decode()
            parts.append(f"<h2>{e(b[1].title)}</h2><figure><img alt='{e(b[1].title)}' src='data:image/png;base64,{img}'><figcaption>{e(b[3])}</figcaption></figure>")
    if doc.notes:
        parts.append("<h2>Data quality</h2><ul>" + "".join(f"<li>{e(n)}</li>" for n in doc.notes) + "</ul>")
    how = f"written by {e(llm.MODEL)}" if doc.ai else "written with built-in rules (the AI model was unavailable)"
    parts.append(f"<footer>Generated by AutoReport AI on {datetime.now():%d %b %Y}. All figures are calculated from the cleaned data; the narrative is {how}.</footer>")
    return f"<!doctype html><html lang='en'><meta charset='utf-8'><meta name='viewport' content='width=device-width,initial-scale=1'><title>{e(doc.title)}</title><style>{_CSS}</style><body><main>{''.join(parts)}</main></body></html>"